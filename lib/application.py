"""
Her yeni ilan icin ana CV'den ilana uyarlanmis bir CV ve on yazi (cover letter) uretir, PDF'e cevirir.

Repo herkese acik oldugu icin:
  - Ana CV repoda sifreli durur (cv/master_cv.docx.enc); workflow CV_KEY secret'i ile /tmp'ye acar.
  - CV icerigi ve uretilen metinler ASLA print edilmez (Actions loglari da herkese acik).

Uydurma korumasi: model sadece ozet, madde (bullet) ve beceri satirlarini degistirebilir. Basliklar, isimler,
tarihler, iletisim bilgileri hic modele birakilmaz. Degistirilen metinde ana CV'de gecmeyen bir ozel isim,
kisaltma veya sayi varsa o degisiklik reddedilir ve orijinal metin kalir. Beceri satirlarinda sadece siralama
degisebilir.
"""
import difflib
import json
import os
import re
import subprocess
import tempfile
from datetime import date

from docx import Document
from docx.shared import Pt, Inches
from openai import OpenAI

from lib.config import NVIDIA_API_KEY
from scrapers.company_nim_fallback import NIM_BASE_URL, NIM_MODEL, _parse_json

MASTER_CV_PATH = os.environ.get("MASTER_CV_PATH", "/tmp/master_cv.docx")
MAX_LENGTH_GROWTH = 1.3  # tek sayfaya sigsin diye: yeni metin orijinalin en fazla %30 uzun olabilir

SYSTEM_PROMPT = """You lightly tailor a student's CV and write a short cover letter for a specific internship posting.
Rules — follow them strictly:
- NEVER invent experience, employers, tools, technologies, certificates, numbers or results that are not in the CV
  or in the student's own note.
- CV: the CV must still read as the student's own. Change ONLY the summary and at most the 4 bullets most relevant to
  this posting. Keep their wording: swap or add a few words so the posting's terms appear where they truthfully fit,
  do not rewrite whole sentences. Leave every other bullet out of the answer. Keep lengths about the same.
- "skills": reorder the comma-separated items of each line so the most relevant come first. Do not add or remove items.
- Cover letter: English, 3 short paragraphs, 150-200 words in total, no greeting and no signature (added separately).
  Write like a real 2nd-year industrial engineering student: plain, specific, short sentences, not marketing language.
  Paragraph 1: which role and one concrete reason it fits (a task from the posting). Paragraph 2: two concrete things
  from the CV that match two tasks in the posting. Paragraph 3: availability for a full-time Erasmus+ internship and a
  simple closing line.
  Do NOT use: em dashes, "I am writing to express", "excited", "thrilled", "passion", "passionate", "leverage", "delve",
  "dynamic", "fast-paced", "I am confident", "unique opportunity", "perfect fit", "align", "invaluable", "honed",
  "spearhead", "testament", "cutting-edge", "furthermore", "moreover".
- "requirements": list the posting's 3-8 most important requirements (short, as the posting states them, e.g. "German
  C1", "SAP experience", "studies in industrial engineering"). Set "met": true ONLY if the CV or the student's note
  clearly shows it; otherwise false.
Reply ONLY with JSON:
{"summary": "...", "bullets": {"<id>": "..."}, "skills": {"<id>": "item, item, ..."}, "cover_letter": ["...", "...", "..."],
 "requirements": [{"req": "...", "met": true}]}"""

# Asiri yeniden yazimi engellemek icin: CV'de en fazla bu kadar madde degisebilir ve degisen metin orijinaline en az
# bu oranda benzemeli (kelime bazinda). Boylece CV her ilanda Kaan'in kendi CV'si gibi okunur.
MAX_BULLET_CHANGES = 4
MIN_SIMILARITY = 0.6

# Okuyana "yapay zeka yazmis" dedirten kalip ifadeler: bunlari iceren cumleler on yazidan cikarilir
AI_CLICHES = (
    "writing to express", "excited", "thrilled", "passion", "leverag", "delve", "dynamic", "fast-paced",
    "i am confident", "i'm confident", "unique opportunity", "perfect fit", "align", "invaluable", "honed",
    "spearhead", "testament", "cutting-edge", "furthermore", "moreover", "in today's", "wealth of",
)

# Ogrencinin kendi yazdigi kisa not (neden staj, neden Avrupa vb.); APPLICANT_NOTE secret'i ile verilir, istege bagli
APPLICANT_NOTE = os.environ.get("APPLICANT_NOTE", "").strip()


# ---------- ana CV'yi okuma ----------

def _is_heading(p) -> bool:
    t = p.text.strip()
    return bool(t) and t.isupper() and len(p.runs) == 1 and bool(p.runs[0].bold)


def _is_bullet(p) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.numPr is not None


def _sections(doc: Document):
    """(bolum_basligi, paragraf_indeksi, paragraf) uclusu dondurur."""
    section = ""
    for i, p in enumerate(doc.paragraphs):
        if _is_heading(p):
            section = p.text.strip()
            continue
        yield section, i, p


def _editable_parts(doc: Document) -> dict:
    parts = {"summary": None, "bullets": {}, "skills": {}}
    for section, i, p in _sections(doc):
        if section == "PROFESSIONAL SUMMARY" and p.text.strip() and parts["summary"] is None:
            parts["summary"] = (i, p.text)
        elif _is_bullet(p) and section != "CERTIFICATIONS" and len(p.runs) == 1:
            parts["bullets"][str(i)] = p.text
        elif section == "TECHNICAL SKILLS" and len(p.runs) >= 2 and p.runs[0].text.strip().endswith(":") \
                and not p.runs[0].text.lower().startswith("language"):
            parts["skills"][str(i)] = p.runs[1].text.strip()
    return parts


def available() -> bool:
    return bool(NVIDIA_API_KEY) and os.path.exists(MASTER_CV_PATH)


# ---------- uydurma korumasi ----------

_TOKEN = re.compile(r"[A-Za-z0-9][\w+#&/.-]*")


def _suspicious_tokens(text: str, corpus_lower: str) -> list[str]:
    """Cumle basi disinda buyuk harfle baslayan, kisaltma veya rakam iceren ve corpus'ta gecmeyen kelimeler."""
    bad = []
    for sentence in re.split(r"(?<=[.;:!?])\s+", text):
        for k, tok in enumerate(_TOKEN.findall(sentence)):
            tok = tok.rstrip(".,")
            special = any(c.isdigit() for c in tok) or (k > 0 and tok[:1].isupper()) or \
                (len(tok) > 1 and tok.isupper())
            if special and tok.lower() not in corpus_lower:
                bad.append(tok)
    return bad


def _similarity(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a.lower().split(), b.lower().split()).ratio()


def _accept(new: str, old: str, corpus_lower: str, min_similarity: float = MIN_SIMILARITY) -> bool:
    if not new or len(new) > len(old) * MAX_LENGTH_GROWTH + 20:
        return False
    if _similarity(new, old) < min_similarity:  # cumle bastan yazilmis: orijinal kalsin
        return False
    return not _suspicious_tokens(new, corpus_lower)


# ---------- model ----------

def _ask_model(parts: dict, offer: dict, job_text: str) -> dict:
    client = OpenAI(base_url=NIM_BASE_URL, api_key=NVIDIA_API_KEY)
    payload = {
        "summary": parts["summary"][1] if parts["summary"] else "",
        "bullets": parts["bullets"],
        "skills": parts["skills"],
    }
    user = (f"POSTING: {offer.get('title')} at {offer.get('company')} ({offer.get('location')})\n"
            f"POSTING TEXT:\n{job_text[:6000]}\n\nEDITABLE CV PARTS (JSON):\n{json.dumps(payload, ensure_ascii=False)}")
    if APPLICANT_NOTE:
        user += f"\n\nSTUDENT'S OWN NOTE (use its facts and tone in the cover letter; never contradict it):\n{APPLICANT_NOTE}"
    for attempt in range(2):
        response = client.chat.completions.create(
            model=NIM_MODEL,
            messages=[{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": user}],
            temperature=0.2,
            max_tokens=3000,
            response_format={"type": "json_object"},
            extra_body={"chat_template_kwargs": {"enable_thinking": False}},
        )
        try:
            return _parse_json(response.choices[0].message.content or "")
        except ValueError:
            if attempt == 1:
                raise


# ---------- belge uretimi ----------

def _set_text(p, text: str, run_index: int = 0):
    p.runs[run_index].text = text


def _build_cv(doc: Document, parts: dict, answer: dict, corpus_lower: str) -> int:
    changed = 0
    if parts["summary"] and isinstance(answer.get("summary"), str):
        i, old = parts["summary"]
        if _accept(answer["summary"].strip(), old, corpus_lower, min_similarity=0.5):
            _set_text(doc.paragraphs[i], answer["summary"].strip())
            changed += 1
    bullets_changed = 0
    for pid, new in (answer.get("bullets") or {}).items():
        if bullets_changed >= MAX_BULLET_CHANGES:
            break
        old = parts["bullets"].get(str(pid))
        if old and isinstance(new, str) and new.strip() != old.strip() and _accept(new.strip(), old, corpus_lower):
            _set_text(doc.paragraphs[int(pid)], new.strip())
            changed += 1
            bullets_changed += 1
    for pid, new in (answer.get("skills") or {}).items():
        old = parts["skills"].get(str(pid))
        if not old or not isinstance(new, str):
            continue
        old_items = [s.strip() for s in old.split(",")]
        new_items = [s.strip() for s in new.split(",")]
        if sorted(old_items) == sorted(new_items):  # sadece siralama degisebilir
            run = doc.paragraphs[int(pid)].runs[1]
            leading = run.text[:len(run.text) - len(run.text.lstrip())]
            run.text = leading + ", ".join(new_items)
            changed += 1
    return changed


def _build_cover_letter(master: Document, paragraphs: list[str], offer: dict, path: str):
    name = master.paragraphs[0].text.strip().title()
    contact = master.paragraphs[2].text.strip()
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    head = doc.add_paragraph()
    r = head.add_run(name)
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph(contact)
    doc.add_paragraph(date.today().strftime("%d %B %Y"))
    doc.add_paragraph(f"Application: {offer.get('title', 'Internship')}, {offer.get('company', '')}")
    doc.add_paragraph(f"Dear Hiring Team at {offer.get('company') or 'your company'},")
    for text in paragraphs:
        doc.add_paragraph(text.strip())
    doc.add_paragraph("Kind regards,")
    doc.add_paragraph(name)
    doc.save(path)


def _to_pdf(docx_path: str, outdir: str) -> str:
    env = dict(os.environ, HOME=tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=True, capture_output=True, timeout=180, env=env)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    if not os.path.exists(pdf):
        raise RuntimeError("PDF olusturulamadi")
    return pdf


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s or "").strip("_")[:40] or "Company"


def build_application(offer: dict, job_text: str) -> dict:
    """{"cv": pdf_yolu, "cover_letter": pdf_yolu (yoksa None), "changes": int} dondurur."""
    master = Document(MASTER_CV_PATH)
    parts = _editable_parts(master)
    cv_text = "\n".join(p.text for p in master.paragraphs)
    answer = _ask_model(parts, offer, job_text)

    outdir = tempfile.mkdtemp(prefix="application_")
    company = _slug(offer.get("company"))

    doc = Document(MASTER_CV_PATH)
    changes = _build_cv(doc, parts, answer, cv_text.lower())
    cv_docx = os.path.join(outdir, f"Kaan_Dogru_CV_{company}.docx")
    doc.save(cv_docx)
    result = {"cv": _to_pdf(cv_docx, outdir), "cover_letter": None, "changes": changes,
              "match": _match_score(answer.get("requirements"))}

    letter_corpus = f"{cv_text}\n{APPLICANT_NOTE}\n{offer.get('title', '')}\n{offer.get('company', '')}\n{job_text}".lower()
    letter, removed = _clean_letter(answer.get("cover_letter") or [], letter_corpus)
    result["letter_sentences_removed"] = removed
    if len(letter) >= 2:
        cl_docx = os.path.join(outdir, f"Kaan_Dogru_Cover_Letter_{company}.docx")
        _build_cover_letter(master, letter, offer, cl_docx)
        result["cover_letter"] = _to_pdf(cl_docx, outdir)
    return result


def _clean_letter(paragraphs: list, corpus_lower: str) -> tuple[list[str], int]:
    """CV'de / ilanda gecmeyen ozel isim, kisaltma veya sayi iceren ve yapay zeka kalibi ifade iceren cumleleri
    cikarir, uzun tireleri virgule cevirir. (temizlenmis_paragraflar, cikarilan_cumle_sayisi) dondurur."""
    cleaned, removed = [], 0
    for p in paragraphs:
        if not isinstance(p, str) or not p.strip():
            continue
        kept = []
        for sentence in re.split(r"(?<=[.!?])\s+", re.sub(r"\s*[—–]\s*", ", ", p.strip())):
            if _suspicious_tokens(sentence, corpus_lower) or any(c in sentence.lower() for c in AI_CLICHES):
                removed += 1
            else:
                kept.append(sentence)
        if kept:
            cleaned.append(" ".join(kept))
    return cleaned, removed


def _match_score(requirements) -> dict | None:
    """Ilanin temel sartlarindan kacini CV'nin karsiladigi. Kabul olasiligi DEGIL, uygunluk tahmini.
    {"score": yuzde, "met": n, "total": n, "missing": [...]} veya (en az 3 sart yoksa) None."""
    reqs = [r for r in (requirements or []) if isinstance(r, dict) and str(r.get("req") or "").strip()][:8]
    if len(reqs) < 3:
        return None
    met = sum(1 for r in reqs if r.get("met") is True)
    missing = [str(r["req"]).strip()[:40] for r in reqs if r.get("met") is not True][:4]
    return {"score": round(100 * met / len(reqs)), "met": met, "total": len(reqs), "missing": missing}
